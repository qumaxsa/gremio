#!/usr/bin/env python3
"""
Genera la LISTA GREMIO Qumax (Hikvision) a partir del paquete JSON del proveedor
(armado por tools/api_proveedor.py desde la API).

Uso:  python3 -I generar_lista_gremio.py <paquete.json> <carpeta_salida>

Entrada (paquete armado en el navegador, ver PROCEDIMIENTO):
  items     -> artículos del proveedor filtrados por Hikvision
  transito  -> in_transit_products.json (clave: item_code = Codigo EAN)
  imgs      -> {Codigo EAN: jpg base64 80x60}

Reglas:
  Precio gremio USD con IVA = Precio Gremio USD (sin IVA) x (1 + IVA%), redondeado a 2 decimales.
  Stock: 0 = Sin stock | 1-9 = Stock bajo | 10-200 = Stock normal | >200 = Stock alto
  Si el articulo tiene mercaderia en transito -> "Ingresando pronto".
  Columnas: Foto, Marca, Articulo, Detalle, Categoria, Subcategoria, Disponibilidad, Ingresos,
            Precio gremio USD (IVA incl.), IVA %.
  Nada del proveedor: sin codigo EAN, sin precio distribuidor, sin links, logos, mails ni telefonos.
"""
import base64, io, json, os, re, sys, html
from collections import Counter, OrderedDict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

# ---------- datos de Qumax (unica info de contacto que aparece) ----------
QUMAX = {
    "nombre": "QUMAX S.A.",
    "web": "www.qumax.com.ar",
    "whatsapp": "+54 9 3492 20-9925",
    "direccion": "Lisandro de la Torre 323, Rafaela, Santa Fe",
    # aplicacion web de Google Apps Script (info@qumax.com.ar) que recibe los pedidos y manda los mails
    "pedidos_url": "https://script.google.com/macros/s/AKfycbyDf0wVkVnxrvMzOM1n_Nsy72fAfsyw5bZK5P2NNJgxxQLG8ywScaiVtnXojtkc8q7uiA/exec",
}
# el nombre del proveedor no se escribe en claro (repo publico): llega por variable o se decodifica
_PROV = [w for w in (os.environ.get("PROV_NOMBRE") or base64.b64decode("Zmllc2E=").decode()).split(",") if w]
FORBIDDEN = re.compile("|".join(map(re.escape, _PROV)) + r"|wa\.me|https?://|www\.(?!qumax)|contacto@|5628|56281711", re.I)

# paleta Qumax: celeste de marca, blanco y grises (sin negro ni rojo)
C_CELESTE, C_CELESTE_OSC, C_CELESTE_CLARO, C_ALT, C_TEXTO, C_GRIS = "0094DD", "0077B3", "E6F4FB", "F4FAFE", "3D4752", "6B7280"
STOCK_COLORS = {"Stock alto": "005E8C", "Stock normal": "0094DD", "Stock bajo": "6B7280", "Sin stock": "A0A7B0"}


CTRL = re.compile(r"[\x00-\x08\x0b-\x1f\x7f]")


def clean(s):
    return re.sub(r"\s+", " ", CTRL.sub(" ", str(s or ""))).strip()


def num(s):
    try:
        return float(str(s).replace(",", ".").strip() or 0)
    except ValueError:
        return 0.0


def stock_label(q):
    if q <= 0:
        return "Sin stock"
    if q < 10:
        return "Stock bajo"
    if q <= 200:
        return "Stock normal"
    return "Stock alto"


def build_rows(pkg):
    transit = {str(p["item_code"]): p for p in pkg["transito"].get("products", [])}
    rows = []
    for x in pkg["items"]:
        if not x.get("Id") or str(x.get("Marca", "")).strip().lower() != "hikvision":
            continue
        ean = str(x["Código EAN"]).strip()
        iva = num(x["IVA"])
        neto = num(x["Precio Gremio USD"])
        if neto <= 0:
            continue
        pods = [p for p in transit.get(ean, {}).get("pods", []) if num(p.get("avalible_quantity", 0)) > 0]
        en_transito = bool(pods)
        llegada = min((datetime.strptime(p["poh_warehouse_est_arrival_date"], "%d/%m/%Y") for p in pods
                       if re.match(r"\d{2}/\d{2}/\d{4}$", str(p.get("poh_warehouse_est_arrival_date", "")))), default=None)
        rows.append(OrderedDict(
            ean=ean,  # solo interno, para la foto; no se exporta
            marca="Hikvision",
            articulo=clean(x["Artículo"]).rstrip("@").strip(),
            detalle=clean(x["Detalle Categoría"]).rstrip(","),
            categoria=clean(x["Categoría"]),
            subcategoria=clean(x["Sub Categoría"]),
            disponibilidad=stock_label(num(x["Stock (*)"])),
            ingresos=(f"Ingresa aprox. {llegada:%d/%m}" if llegada else "Ingresando pronto") if en_transito else "",
            precio=round(neto * (1 + iva / 100), 2),
            iva=iva,
        ))
    rows.sort(key=lambda r: (r["categoria"], r["subcategoria"], r["articulo"]))
    return rows


def check_clean(rows):
    bad = [r["articulo"] for r in rows if any(FORBIDDEN.search(str(r[k])) for k in r if k != "ean")]
    if bad:
        raise SystemExit(f"ERROR: referencias prohibidas en {bad[:10]}")


# ------------------------------- EXCEL -------------------------------
HEAD = ["Foto", "Marca", "Artículo", "Detalle", "Categoría", "Subcategoría",
        "Disponibilidad", "Ingresos", "Precio gremio USD\n(IVA incluido)", "IVA %"]
WIDTHS = [13, 11, 24, 58, 18, 22, 15, 17, 17, 8]


def write_sheet(ws, rows, imgs, fecha, titulo):
    thin = Side(style="thin", color="D6E9F5")
    ncol = len(HEAD)
    last = get_column_letter(ncol)
    ws.merge_cells(f"A1:{last}1"); ws.merge_cells(f"A2:{last}2"); ws.merge_cells(f"A3:{last}3")
    ws["A1"] = f"QUMAX  ·  LISTA GREMIO HIKVISION  ·  {titulo}"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    ws["A2"] = (f"Actualizada el {fecha}  ·  Precios de referencia para instaladores, en dólares con IVA incluido. "
                "Sujetos a cambio sin previo aviso. Disponibilidad aproximada.")
    ws["A2"].font = Font(size=10, color="FFFFFF")
    ws["A3"] = f"Pedidos: WhatsApp {QUMAX['whatsapp']}  ·  {QUMAX['web']}  ·  {QUMAX['direccion']}"
    ws["A3"].font = Font(size=10, bold=True, color=C_CELESTE_OSC)
    for r, fill in ((1, C_CELESTE), (2, C_CELESTE), (3, C_CELESTE_CLARO)):
        for c in range(1, ncol + 1):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=fill)
        ws.cell(r, 1).alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 18; ws.row_dimensions[3].height = 20

    hr = 5
    for c, h in enumerate(HEAD, 1):
        cell = ws.cell(hr, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=C_CELESTE_OSC)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.row_dimensions[hr].height = 32

    for i, r in enumerate(rows):
        rr = hr + 1 + i
        vals = ["", r["marca"], r["articulo"], r["detalle"], r["categoria"], r["subcategoria"],
                r["disponibilidad"], r["ingresos"], r["precio"], r["iva"] / 100]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(rr, c, v)
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 4))
            cell.border = Border(bottom=thin)
            cell.font = Font(color=C_TEXTO)
            if i % 2:
                cell.fill = PatternFill("solid", fgColor=C_ALT)
        ws.cell(rr, 3).font = Font(bold=True, color=C_TEXTO)
        ws.cell(rr, 7).font = Font(bold=True, color=STOCK_COLORS[r["disponibilidad"]])
        ws.cell(rr, 8).font = Font(bold=True, color=C_CELESTE_OSC)
        ws.cell(rr, 9).number_format = '"US$" #,##0.00'
        ws.cell(rr, 9).font = Font(bold=True, size=11, color=C_CELESTE_OSC)
        ws.cell(rr, 10).number_format = "0.0%"
        ws.cell(rr, 10).font = Font(color=C_GRIS)
        for c in (7, 8, 9, 10):
            ws.cell(rr, c).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[rr].height = 48
        b64 = imgs.get(r["ean"])
        if b64:
            img = XLImage(io.BytesIO(base64.b64decode(b64)))
            img.width, img.height = 80, 60
            ws.add_image(img, f"A{rr}")

    ws.freeze_panes = ws.cell(hr + 1, 3)
    ws.auto_filter.ref = f"A{hr}:{last}{hr + len(rows)}"
    ws.print_title_rows = f"{hr}:{hr}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)
    ws.oddFooter.center.text = "QUMAX S.A. · Lista gremio · Página &P de &N"


def build_excel(rows, imgs, fecha, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Todas las categorías"
    write_sheet(ws, rows, imgs, fecha, "TODAS LAS CATEGORÍAS")
    for cat in sorted({r["categoria"] for r in rows}):
        sub = [r for r in rows if r["categoria"] == cat]
        write_sheet(wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "", cat)[:31]), sub, imgs, fecha, cat.upper())
    wb.properties.creator = "Qumax S.A."
    wb.properties.title = "Lista gremio Hikvision - Qumax"
    wb.properties.company = "Qumax S.A."
    wb.save(out)


# ------------------------------- WEB -------------------------------
def visibles(rows):
    """En la web, al navegar, no se muestran los artículos sin stock salvo que estén ingresando pronto.
    Si el técnico escribe algo en el buscador, aparecen todos (lo resuelve la página; regla del 29/10/2026)."""
    return [r for r in rows if r["disponibilidad"] != "Sin stock" or r["ingresos"]]


def build_html(rows, imgs, fecha, out):
    data = [{k: r[k] for k in ("articulo", "detalle", "categoria", "subcategoria", "disponibilidad", "ingresos", "precio", "iva")}
            | {"img": ("data:image/jpeg;base64," + imgs[r["ean"]]) if r["ean"] in imgs else ""} for r in rows]
    tpl = Path(__file__).with_name("plantilla_web.html").read_text(encoding="utf-8")
    page = (tpl.replace("__FECHA__", html.escape(fecha))
               .replace("__WHATSAPP__", QUMAX["whatsapp"])
               .replace("__WA_NUM__", re.sub(r"\D", "", QUMAX["whatsapp"]))
               .replace("__DATA__", json.dumps(data, ensure_ascii=False).replace("</", "<\\/")))
    Path(out).write_text(page, encoding="utf-8")
    # pagina de ingreso del sitio con clave (sin datos: los datos van cifrados aparte)
    priv = Path(__file__).with_name("plantilla_gremio_privado.html").read_text(encoding="utf-8")
    Path(out).with_name("index.html").write_text(
        priv.replace("__WHATSAPP__", QUMAX["whatsapp"]).replace("__WA_NUM__", re.sub(r"\D", "", QUMAX["whatsapp"]))
            .replace("__PEDIDOS_URL__", sys.argv[3] if len(sys.argv) > 3 else QUMAX["pedidos_url"]),
        encoding="utf-8")
    # datos para el sitio con clave (se cifran con gremio_crypto.py cifrar)
    Path(out).with_name(Path(out).stem.replace("Lista_Gremio_Qumax_Hikvision", "web_data") + ".json").write_text(
        json.dumps({"fecha": fecha, "items": data}, ensure_ascii=False), encoding="utf-8")


def main():
    src, outdir = Path(sys.argv[1]), Path(sys.argv[2])
    outdir.mkdir(parents=True, exist_ok=True)
    pkg = json.loads(src.read_text(encoding="utf-8"))
    fecha_dt = datetime.fromisoformat(pkg["fecha"].replace("Z", "+00:00")).astimezone(ZoneInfo("America/Argentina/Buenos_Aires"))
    fecha = fecha_dt.strftime("%d/%m/%Y %H:%M") if os.environ.get("GREMIO_CON_HORA") else fecha_dt.strftime("%d/%m/%Y")
    tag = fecha_dt.strftime("%Y-%m-%d")
    rows = build_rows(pkg)
    check_clean(rows)
    imgs = pkg.get("imgs", {})
    # la imagen generica de "sin foto" se repite muchas veces: se descarta
    reps = Counter(imgs.values())
    imgs = {k: v for k, v in imgs.items() if reps[v] <= 20}
    xlsx = outdir / f"Lista_Gremio_Qumax_Hikvision_{tag}.xlsx"
    web = outdir / f"Lista_Gremio_Qumax_Hikvision_{tag}.html"
    if not os.environ.get("GREMIO_SIN_EXCEL"):
        build_excel(rows, imgs, fecha, xlsx)
    build_html(rows, imgs, fecha, web)
    json.dump([{k: v for k, v in r.items() if k != "ean"} for r in rows],
              open(outdir / f"lista_gremio_{tag}.json", "w", encoding="utf-8"), ensure_ascii=False)
    c = Counter(r["disponibilidad"] for r in rows)
    print(f"Web: {len(visibles(rows))} artículos visibles al navegar ({len(rows) - len(visibles(rows))} sin stock y sin ingreso, solo aparecen al buscarlos)")
    print(f"OK {len(rows)} artículos | con foto {sum(r['ean'] in imgs for r in rows)} | "
          f"ingresando pronto {sum(bool(r['ingresos']) for r in rows)} | {dict(c)}")
    print(xlsx); print(web)


if __name__ == "__main__":
    main()
